"""Generate HubSpot-format bulk-upload XLSX files for social posts.

HubSpot's social bulk-upload sheet expects the columns
Account, Date, Message, Link, Image URL, with dates as MM/DD/YY HH:MM in the
portal's timezone.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import pendulum
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

HEADERS = ["Account", "Date", "Message", "Link", "Image URL"]
MAX_POSTS = 300  # HubSpot's documented bulk-upload ceiling

# Rough platform ceilings, used for warnings only — never to reject a post.
PLATFORM_LIMITS = {
    "linkedin": 3000,
    "facebook": 63206,
    "instagram": 2200,
    "x": 280,
    "twitter": 280,
}


# Excel evaluates a cell starting with any of these as a formula, including
# the =cmd|...!A1 DDE form that reaches code execution. Every text cell we
# write may contain model-supplied content, so every text cell is quoted.
_FORMULA_LEAD = ("=", "+", "-", "@", "\t", "\r")

# openpyxl rejects these outright, and they turn up routinely in text copied
# out of HubSpot rich-text fields.
_ILLEGAL_XLSX = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

MAX_CELL_CHARS = 32_767  # Excel's hard ceiling


class SocialXLSXError(ValueError):
    """Raised for any input the caller can fix, with a message that says how."""


def _as_text(value: Any) -> str:
    """Render a value as an inert Excel string."""
    text = "" if value is None else str(value)
    text = _ILLEGAL_XLSX.sub("", text)
    if len(text) > MAX_CELL_CHARS:
        text = text[: MAX_CELL_CHARS - 1]
    if text.startswith(_FORMULA_LEAD):
        text = "'" + text
    return text


def _parse_dt(raw: Any, tz: str, *, row: int) -> pendulum.DateTime:
    if not isinstance(raw, str):
        raise SocialXLSXError(
            f"Post #{row}: scheduled_at must be an ISO 8601 string, got {type(raw).__name__}."
        )
    try:
        parsed = pendulum.parse(raw, tz=tz)
    except Exception as exc:
        raise SocialXLSXError(
            f"Post #{row}: could not read scheduled_at={raw!r}. "
            f"Use ISO 8601, e.g. '2026-01-13T09:00:00'. ({exc})"
        ) from None
    if not isinstance(parsed, pendulum.DateTime):
        raise SocialXLSXError(f"Post #{row}: scheduled_at={raw!r} is not a date and time.")
    return parsed.in_timezone(tz)


def _platform_warning(account: str, message: str, row: int) -> str | None:
    lowered = account.lower()
    for platform, limit in PLATFORM_LIMITS.items():
        if platform in lowered and len(message) > limit:
            return (
                f"Post #{row} ({account}): {len(message)} characters exceeds the "
                f"~{limit} character limit for {platform}."
            )
    return None


def generate_social_bulk_xlsx(
    posts: list[dict[str, Any]],
    *,
    output_dir: Path,
    output_filename: str | None = None,
    timezone: str = "UTC",
) -> dict[str, Any]:
    if not posts:
        raise SocialXLSXError("posts is empty — supply at least one post.")
    if len(posts) > MAX_POSTS:
        raise SocialXLSXError(
            f"HubSpot accepts at most {MAX_POSTS} posts per bulk-upload file; "
            f"got {len(posts)}. Split them across several files."
        )

    # Fail on a bad timezone once, before touching any row.
    try:
        pendulum.timezone(timezone)
    except Exception as exc:
        raise SocialXLSXError(
            f"Unknown timezone {timezone!r}. Use an IANA name such as "
            f"'Europe/Berlin' or 'America/New_York'. ({exc})"
        ) from None

    wb = Workbook()
    ws = wb.active
    ws.title = "Posts"

    bold = Font(bold=True)
    for col, header in enumerate(HEADERS, start=1):
        ws.cell(row=1, column=col, value=header).font = bold

    accounts: set[str] = set()
    moments: list[pendulum.DateTime] = []
    warnings: list[str] = []
    wrap = Alignment(wrap_text=True, vertical="top")

    for offset, post in enumerate(posts):
        row_no = offset + 1
        sheet_row = offset + 2

        if not isinstance(post, dict):
            raise SocialXLSXError(
                f"Post #{row_no} must be an object with account, scheduled_at and message."
            )
        missing = [k for k in ("account", "scheduled_at", "message") if not post.get(k)]
        if missing:
            raise SocialXLSXError(
                f"Post #{row_no} is missing required field(s): {', '.join(missing)}."
            )

        account = str(post["account"])
        message = str(post["message"])
        moment = _parse_dt(post["scheduled_at"], timezone, row=row_no)

        warning = _platform_warning(account, message, row_no)
        if warning:
            warnings.append(warning)

        accounts.add(account)
        moments.append(moment)

        ws.cell(row=sheet_row, column=1, value=_as_text(account))
        ws.cell(row=sheet_row, column=2, value=moment.format("MM/DD/YY HH:mm"))
        ws.cell(row=sheet_row, column=3, value=_as_text(message)).alignment = wrap
        ws.cell(row=sheet_row, column=4, value=_as_text(post.get("link_url")))
        ws.cell(row=sheet_row, column=5, value=_as_text(post.get("image_url")))

    for col, width in enumerate([28, 16, 80, 40, 40], start=1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width

    output_dir.mkdir(parents=True, exist_ok=True)

    if output_filename:
        # Only ever a bare filename — never let a caller escape output_dir.
        safe_name = Path(str(output_filename)).name
        if not safe_name or safe_name in {".", ".."}:
            raise SocialXLSXError(f"Invalid output_filename: {output_filename!r}")
    else:
        safe_name = f"social_schedule_{pendulum.now(timezone).format('YYYY-MM-DD')}.xlsx"
    if not safe_name.endswith(".xlsx"):
        safe_name += ".xlsx"

    file_path = output_dir / safe_name
    if file_path.is_symlink():
        raise SocialXLSXError(
            f"Refusing to write through a symlink at {file_path}. Remove it and retry."
        )
    try:
        wb.save(file_path)
    except OSError as exc:
        raise SocialXLSXError(
            f"Could not write {file_path}: {exc}. Check that the output directory "
            f"exists and is writable."
        ) from None

    moments.sort()
    result: dict[str, Any] = {
        "file_path": str(file_path),
        "post_count": len(posts),
        "accounts": sorted(accounts),
        "timezone": timezone,
        "date_range": {
            "start": moments[0].to_iso8601_string(),
            "end": moments[-1].to_iso8601_string(),
        },
        "next_step": (
            "In HubSpot go to Marketing → Social → 'Schedule in bulk' and upload "
            "this file. HubSpot previews each post before anything is scheduled."
        ),
    }
    if warnings:
        result["warnings"] = warnings
    return result
