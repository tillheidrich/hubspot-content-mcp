from __future__ import annotations

import pytest
from openpyxl import load_workbook

from hubspot_mcp.social.xlsx_writer import (
    HEADERS,
    MAX_POSTS,
    SocialXLSXError,
    generate_social_bulk_xlsx,
)


def _post(**overrides):
    base = {
        "account": "Acme Inc - LinkedIn",
        "scheduled_at": "2026-06-09T09:00:00",
        "message": "Hello world",
    }
    base.update(overrides)
    return base


def test_writes_the_hubspot_column_layout(tmp_path):
    out = generate_social_bulk_xlsx([_post()], output_dir=tmp_path)

    wb = load_workbook(out["file_path"])
    ws = wb.active
    assert ws.title == "Posts"
    assert [c.value for c in ws[1]] == HEADERS
    assert ws["A2"].value == "Acme Inc - LinkedIn"
    assert ws["C2"].value == "Hello world"


def test_dates_use_hubspots_mm_dd_yy_format(tmp_path):
    out = generate_social_bulk_xlsx(
        [_post(scheduled_at="2026-06-09T09:30:00")],
        output_dir=tmp_path,
        timezone="Europe/Berlin",
    )
    ws = load_workbook(out["file_path"]).active
    assert ws["B2"].value == "06/09/26 09:30"


def test_date_range_is_chronological_not_lexicographic(tmp_path):
    """Mixed offsets sort wrong as strings."""
    out = generate_social_bulk_xlsx(
        [
            _post(scheduled_at="2026-06-09T09:00:00+02:00"),
            _post(scheduled_at="2026-06-09T06:30:00Z"),  # earlier in real time
        ],
        output_dir=tmp_path,
        timezone="Europe/Berlin",
    )
    assert out["date_range"]["start"].startswith("2026-06-09T08:30")


def test_empty_post_list_is_rejected(tmp_path):
    with pytest.raises(SocialXLSXError, match="empty"):
        generate_social_bulk_xlsx([], output_dir=tmp_path)


def test_too_many_posts_is_rejected(tmp_path):
    with pytest.raises(SocialXLSXError, match=str(MAX_POSTS)):
        generate_social_bulk_xlsx([_post()] * (MAX_POSTS + 1), output_dir=tmp_path)


def test_missing_field_names_the_row_and_the_field(tmp_path):
    with pytest.raises(SocialXLSXError) as excinfo:
        generate_social_bulk_xlsx([_post(), {"account": "x", "message": "y"}], output_dir=tmp_path)
    assert "Post #2" in str(excinfo.value)
    assert "scheduled_at" in str(excinfo.value)


def test_unparseable_date_names_the_row_and_suggests_a_format(tmp_path):
    with pytest.raises(SocialXLSXError) as excinfo:
        generate_social_bulk_xlsx([_post(scheduled_at="13.01.2026 09:00")], output_dir=tmp_path)
    message = str(excinfo.value)
    assert "Post #1" in message
    assert "ISO 8601" in message


def test_duration_string_is_not_mistaken_for_a_datetime(tmp_path):
    with pytest.raises(SocialXLSXError, match="not a date"):
        generate_social_bulk_xlsx([_post(scheduled_at="P1D")], output_dir=tmp_path)


def test_bad_timezone_fails_before_any_row_is_processed(tmp_path):
    with pytest.raises(SocialXLSXError, match="[Uu]nknown timezone"):
        generate_social_bulk_xlsx([_post()], output_dir=tmp_path, timezone="Not/AZone")


def test_filename_cannot_escape_the_output_directory(tmp_path):
    out = generate_social_bulk_xlsx(
        [_post()],
        output_dir=tmp_path,
        output_filename="../../../escaped.xlsx",
    )
    written = out["file_path"]
    assert written.startswith(str(tmp_path))
    assert "escaped.xlsx" in written


def test_xlsx_extension_is_added_when_missing(tmp_path):
    out = generate_social_bulk_xlsx([_post()], output_dir=tmp_path, output_filename="week1")
    assert out["file_path"].endswith("week1.xlsx")


def test_overlong_message_produces_a_warning_not_an_error(tmp_path):
    out = generate_social_bulk_xlsx(
        [_post(account="Acme - X", message="x" * 500)], output_dir=tmp_path
    )
    assert out["post_count"] == 1
    assert out["warnings"]
    assert "280" in out["warnings"][0]


def test_result_includes_the_next_step_for_the_user(tmp_path):
    out = generate_social_bulk_xlsx([_post()], output_dir=tmp_path)
    assert "Schedule in bulk" in out["next_step"]
